import asyncio
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import httpx
import openpyxl
from fastapi.testclient import TestClient

import app
import research
from jobs import JobStore, WorkerLock
from research import assemble, evidence, investigate, property_key


ROW = dict(street='12-14 Example St', city='Example City', state='FL', zip_code='32000', row_number=2)


def parcel(owner='Example Holdings LLC'):
    return evidence('official_parcel', 'Mock county', 'https://example.invalid/parcel',
        owner=owner, address_match=True, apn='MOCK-1', mailing_address='Mock mailing address')


class EvidenceTests(unittest.IsolatedAsyncioTestCase):
    def test_registered_agent_and_attorney_are_not_owners(self):
        records = [parcel(), evidence('corporate_filing', 'Mock SOS', entity='Example Holdings LLC',
            person='Example Attorney', role='Incorporation attorney / registered agent')]
        d = assemble(ROW, records)
        self.assertEqual(d['primary_decision_maker'], '')
        self.assertEqual(d['contacts'], [])
        self.assertIn('registered agent', d['principal_candidates'][0]['role'])

    def test_out_of_state_phone_is_retained_and_person_specific(self):
        d = assemble(ROW, [parcel('Example Person'), evidence('phone', 'Mock authorized provider',
            person='Example Person', number='212-555-0100', name_match=True, address_match=True,
            original_record=True, independent_lineages=['db-one','db-two']),
            evidence('phone', 'Mock provider', person='Unrelated Person', number='305-555-0101')])
        self.assertEqual(len(d['phone_candidates']), 2)
        self.assertEqual(d['contacts'][0]['phones'][0]['number'], '212-555-0100')
        self.assertEqual(len(d['contacts'][0]['phones']), 1)
        self.assertEqual(d['primary_phone'], '')
        self.assertNotIn('High', d['overall_confidence'])

    def test_repeated_underlying_database_is_not_independent(self):
        d = assemble(ROW, [evidence('phone', 'Two websites', person='Example Person', number='212-555-0100',
            name_match=True, address_match=True, original_record=True, independent_lineages=['same-db','same-db'])])
        self.assertEqual(d['phone_candidates'][0]['verification_status'], 'unverified')

    def test_conflicting_search_owner_names(self):
        d = assemble(ROW, [evidence('search_snippet', 'Mock search A', owner='Person A'),
            evidence('search_snippet', 'Mock search B', owner='Person B')])
        self.assertEqual(d['owner_of_record'], '')
        self.assertTrue(d['conflicts'])

    def test_conflicting_official_records_withhold_owner(self):
        d = assemble(ROW, [parcel('Person A'), parcel('Person B')])
        self.assertEqual(d['owner_of_record'], '')
        self.assertTrue(d['conflicts'])

    async def test_gemini_fabrications_cannot_override_records(self):
        sources = AsyncMock()
        sources.retrieve.return_value = [parcel()]
        advisor = AsyncMock()
        advisor.advise.return_value = dict(status='completed', owner_of_record='Invented Owner',
            primary_phone='212-555-0100', overall_confidence='High', hypotheses=['Invented Owner is a hypothesis'])
        d = await investigate(ROW, sources, advisor)
        self.assertEqual(d['owner_of_record'], 'Example Holdings LLC')
        self.assertEqual(d['primary_phone'], '')
        self.assertEqual(d['contacts'], [])
        self.assertNotEqual(d['overall_confidence'], 'High')

    async def test_ai_failure_is_visible(self):
        sources, advisor = AsyncMock(), AsyncMock()
        sources.retrieve.return_value = []
        advisor.advise.side_effect = RuntimeError('secret must not be displayed')
        d = await investigate(ROW, sources, advisor)
        self.assertEqual(d['ai']['status'], 'failed')
        self.assertNotIn('secret', json.dumps(d))

    async def test_gemini_uses_supported_structured_output(self):
        from types import SimpleNamespace
        response = SimpleNamespace(text='{"next_steps":["Check the official deed"]}')
        session = AsyncMock()
        session.models.generate_content = AsyncMock(return_value=response)
        client = __import__('unittest').mock.MagicMock()
        client.aio.__aenter__ = AsyncMock(return_value=session)
        client.aio.__aexit__ = AsyncMock(return_value=False)
        with patch('google.genai.Client', return_value=client):
            result = await research.GeminiAdvisor('fake-key','gemini-flash-latest').advise([])
        config = session.models.generate_content.call_args.kwargs['config']
        self.assertEqual(config.response_mime_type, 'application/json')
        self.assertEqual(result['status'], 'completed')


class ParserTests(unittest.TestCase):
    def test_excel_first_sheet_suite_and_hyphenated_street(self):
        wb = openpyxl.Workbook()
        wb.active.append(['Full address','Owner','Tax mailing address'])
        wb.active.append(['12-14 Example St, Suite 200, Example City, FL 32000','Example LLC','Another City, NY'])
        wb.create_sheet('Ignored').append(['Do not process'])
        stream = io.BytesIO()
        wb.save(stream)
        parsed = app.extract_raw_rows_from_bytes(stream.getvalue(), 'mock.xlsx')
        headers, rows = parsed['rows'][0], parsed['rows'][1:]
        columns = app.classify_top_row_headers(headers, rows)
        mapped = app.build_mapped_rows_from_table(headers, rows, [c['mapped_to'] for c in columns])
        self.assertEqual(len(mapped), 1)
        self.assertEqual(mapped[0]['street'], '12-14 Example St, Suite 200')
        self.assertEqual(mapped[0]['city'], 'Example City')
        self.assertEqual(mapped[0]['state'], 'FL')
        self.assertEqual(mapped[0]['mailing_address_hint'], 'Another City, NY')

    def test_missing_city_state_stay_blank(self):
        row = app.build_mapped_rows_from_table(['Street'], [['12 Example St']], ['street'])[0]
        self.assertEqual((row['city'], row['state']), ('',''))
        self.assertIn('city, state', row['validation_error'])
        self.assertEqual(app.normalize_state_abbr(''), '')
        self.assertEqual(app.BatchParseTextRequest(raw_text='').default_state, '')

    def test_mailing_address_is_never_promoted(self):
        columns = app.classify_top_row_headers(['Tax mailing address'], [['12 Example St']])
        self.assertEqual(columns[0]['mapped_to'], 'mailing_address')
        row = app.build_mapped_rows_from_table(['Tax mailing address'], [['12 Example St']], ['mailing_address'])[0]
        self.assertEqual(row['street'], '')
        self.assertTrue(row['validation_error'])

    def test_pasted_table_always_uses_row_one_as_header(self):
        result = asyncio.run(app.parse_batch_text(app.BatchParseTextRequest(raw_text='Address\tCity\tState\n12-14 Example St\tExample City\tFL')))
        self.assertEqual(len(result['mapped_rows']), 1)
        self.assertEqual(result['mapped_rows'][0]['street'], '12-14 Example St')

    def test_word_is_rejected(self):
        with self.assertRaises(app.HTTPException):
            app.extract_raw_rows_from_bytes(b'not a document', 'mock.docx')

    def test_full_state_name_is_normalized_without_invented_state(self):
        self.assertEqual(app.normalize_state_abbr('Florida'), 'FL')
        self.assertNotEqual(app.normalize_state_abbr('incorrect'), 'IN')

    def test_tax_location_columns_never_populate_subject_location(self):
        columns = app.classify_top_row_headers(['Tax City', 'Tax State', 'Tax ZIP'], [['Mock','NC','27601']])
        self.assertEqual([c['mapped_to'] for c in columns], ['notes','notes','notes'])

    def test_csv_quoted_multiline_notes_are_preserved(self):
        rows = app.extract_raw_rows_from_text('Address,Notes\n12 Example St,"First line\nSecond line"')['rows']
        self.assertEqual(rows[1][1], 'First line\nSecond line')


class SourceTests(unittest.IsolatedAsyncioTestCase):
    async def test_rate_limit_retry_after_suppresses_further_requests(self):
        calls = []
        def handler(request):
            calls.append(request)
            return httpx.Response(429, headers={'Retry-After':'120'})
        research.SOURCE_COOLDOWNS.clear()
        try:
            async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
                with self.assertRaises(httpx.HTTPStatusError):
                    await research.public_get(client, 'https://example.invalid/mock')
                with self.assertRaises(ValueError):
                    await research.public_get(client, 'https://example.invalid/mock')
            self.assertEqual(len(calls), 1)
        finally:
            research.SOURCE_COOLDOWNS.clear()

    async def retrieve(self, row, candidates, features=()):
        calls = []
        def handler(request):
            calls.append(request)
            return httpx.Response(200, json={'candidates': candidates} if 'GeocodeServer' in str(request.url)
                else {'features': [{'attributes': f} for f in features]})
        client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
        with patch('research.httpx.AsyncClient', return_value=client):
            records = await research.PublicSources().retrieve(row)
        return records, calls

    def candidate(self, row):
        return dict(score=100, address='Mock address', location={'x':0, 'y':0}, attributes=dict(
            StAddr=row['street'], City=row['city'], RegionAbbr=row['state'], Postal=row['zip_code']))

    async def test_non_nc_state_never_uses_nc_map(self):
        records, calls = await self.retrieve(ROW, [self.candidate(ROW)])
        self.assertEqual(len(calls), 1)
        self.assertEqual(records[-1]['status'], 'not_configured')

    async def test_wrong_parcel_is_not_selected(self):
        row = dict(ROW, state='NC')
        records, calls = await self.retrieve(row, [self.candidate(row)], [dict(siteadd='99 Wrong St', scity=row['city'], ownname='Wrong Owner')])
        self.assertFalse(any(r['kind'] == 'official_parcel' for r in records))
        self.assertEqual(records[-1]['status'], 'manual_review')
        self.assertNotIn('distance', calls[-1].url.params)

    async def test_exact_parcel_has_provenance(self):
        row = dict(ROW, state='NC')
        records, _ = await self.retrieve(row, [self.candidate(row)], [dict(siteadd=row['street'], scity=row['city'], ownname='Mock LLC', parno='1')])
        self.assertEqual(records[-1]['kind'], 'official_parcel')
        self.assertIn('raw_record', records[-1])
        self.assertTrue(records[-1]['source_url'])

    async def test_geocode_mismatch_prevents_parcel_lookup(self):
        row = dict(ROW, state='NC')
        candidate = self.candidate(row)
        candidate['attributes']['RegionAbbr'] = 'NY'
        records, calls = await self.retrieve(row, [candidate])
        self.assertEqual(len(calls), 1)
        self.assertEqual(records[-1]['status'], 'manual_review')


class QueueTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.store = JobStore(Path(self.temp.name) / 'jobs.sqlite3')
        self.saved = {}

    async def asyncTearDown(self):
        self.temp.cleanup()

    async def research(self, row):
        return assemble(row, [parcel()])

    def save(self, result, row):
        self.saved.setdefault(property_key(row), result)

    async def test_duplicates_and_invalid_rows_are_retained_in_export(self):
        job = self.store.create([ROW, dict(ROW, zip_code=''), dict(ROW, street='')])
        await self.store.step(self.research, self.save)
        data = self.store.get(job)
        self.assertEqual([r['status'] for r in data['rows']], ['succeeded','duplicate','invalid'])
        wb = openpyxl.load_workbook(io.BytesIO(self.store.export(job)))
        self.assertEqual(wb.active.max_row, 4)
        self.assertEqual(wb.active.cell(3,12).value, 'duplicate')
        wb.close()

    async def test_cancel_finishes_active_row_then_resume(self):
        job = self.store.create([ROW, dict(ROW, street='22 Other St')])
        async def cancel_during_row(row):
            self.store.action(job, 'cancel')
            return await self.research(row)
        await self.store.step(cancel_during_row, self.save)
        self.assertEqual(self.store.get(job)['status'], 'paused')
        self.assertEqual(self.store.get(job)['counts']['queued'], 1)
        self.store.action(job, 'resume')
        await self.store.step(self.research, self.save)
        self.assertEqual(self.store.get(job)['status'], 'completed')

    async def test_interruption_recovery_and_idempotent_replay(self):
        job = self.store.create([ROW])
        async def save_then_interrupt(row):
            self.save(await self.research(row), row)
            raise asyncio.CancelledError()
        with self.assertRaises(asyncio.CancelledError):
            await self.store.step(save_then_interrupt, self.save)
        store = JobStore(self.store.path)
        store.recover()
        self.assertEqual(store.get(job)['status'], 'paused')
        store.action(job, 'resume')
        await store.step(self.research, self.save)
        self.assertEqual(len(self.saved), 1)
        self.assertEqual(store.get(job)['rows'][0]['attempts'], 2)

    async def test_failed_rows_retry_without_repeating_successes(self):
        job = self.store.create([ROW, dict(ROW, street='22 Other St')])
        await self.store.step(self.research, self.save)
        failing = AsyncMock(side_effect=RuntimeError('mock error'))
        await self.store.step(failing, self.save)
        self.assertEqual(self.store.get(job)['counts']['failed'], 1)
        self.store.action(job, 'retry')
        await self.store.step(self.research, self.save)
        rows = self.store.get(job)['rows']
        self.assertEqual([r['attempts'] for r in rows], [1,2])
        self.assertEqual(len(self.saved), 2)

    async def test_export_does_not_execute_spreadsheet_formulas(self):
        job = self.store.create([dict(ROW, notes='=WEBSERVICE("example.invalid")')])
        wb = openpyxl.load_workbook(io.BytesIO(self.store.export(job)))
        self.assertEqual(wb.active.cell(2,11).data_type, 's')
        wb.close()


class IntegrityTests(unittest.TestCase):
    def test_retries_preserve_saved_record_and_corrupt_file(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(app, 'LEADS_FILE', Path(directory)/'leads.json'):
            payload = dict(address='12 Example St, Example City, FL', notes='manual edit')
            first = app.persist_lead(payload, 'mock-key')
            second = app.persist_lead(dict(payload, notes='retry'), 'mock-key')
            self.assertEqual(first, second)
            self.assertEqual(len(app.load_saved_leads()), 1)
            app.LEADS_FILE.write_text('corrupted', encoding='utf-8')
            with self.assertRaises(ValueError):
                app.persist_lead(payload, 'different-key')
            self.assertEqual(app.LEADS_FILE.read_text(), 'corrupted')

    def test_key_save_preserves_other_configuration(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(app, 'CONFIG_FILE', Path(directory)/'config.json'), patch.dict('os.environ'):
            app.CONFIG_FILE.write_text('{"unrelated_setting":true}')
            app.save_api_key('fake-key')
            self.assertTrue(json.loads(app.CONFIG_FILE.read_text())['unrelated_setting'])

    def test_worker_lock_rejects_second_process(self):
        with tempfile.TemporaryDirectory() as directory:
            first, second = WorkerLock(Path(directory)/'worker.lock'), WorkerLock(Path(directory)/'worker.lock')
            first.acquire()
            try:
                with self.assertRaises(RuntimeError):
                    second.acquire()
            finally:
                first.release()

    def test_job_api_and_cross_origin_guard_with_isolated_storage(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(app, 'BASE_DIR', Path(directory)), patch.object(app, 'JOB_FILE', Path(directory)/'jobs.sqlite3'):
            with TestClient(app.app) as client:
                response = client.post('/api/batch/jobs', json={'rows':[dict(ROW, street='', gemini_api_key='must-not-persist')]})
                self.assertEqual(response.status_code, 200)
                job_id = response.json()['id']
                job = client.get('/api/batch/jobs/' + job_id).json()
                self.assertNotIn('gemini_api_key', job['rows'][0]['payload'])
                self.assertEqual(job['counts']['invalid'], 1)
                self.assertEqual(client.get(f'/api/batch/jobs/{job_id}/export').status_code, 200)
                self.assertEqual(client.post('/api/batch/jobs', json={'rows':[ROW]}, headers={'Origin':'https://example.invalid'}).status_code, 403)
                self.assertEqual(client.get('/api/batch/jobs/missing').status_code, 404)


if __name__ == '__main__':
    unittest.main()
